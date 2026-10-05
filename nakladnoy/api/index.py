"""Nakladnoy bot (alohida bot, Vercel serverless, Flask).

Mijoz yozgan ro'yxatdan (masalan "A10S 10 ta") Xitoy omborining narx faylidagi
modellarni topib, Format1 yoki Format2 ko'rinishidagi Excel nakladnoy tayyorlaydi.

Vercel > Settings > Environment Variables:
  NAKLADNOY_BOT_TOKEN       - @BotFather dan olingan yangi bot tokeni
  NAKLADNOY_ADMIN_IDS       - botdan foydalana oladiganlar Telegram ID lari, vergul bilan: 123,456
  NAKLADNOY_GROUP_IDS       - (ixtiyoriy) barcha a'zolari foydalana oladigan guruh ID lari
Guruhda: botni guruh admini qiling (yoki @BotFather > /setprivacy > Disable), aks holda
bot faqat buyruqlar va o'ziga javob (reply) qilingan xabarlarni ko'radi.
  NAKLADNOY_WEBHOOK_SECRET  - istalgan uzun tasodifiy satr (faqat A-Z a-z 0-9 _ -)
  UPSTASH_REDIS_REST_URL, UPSTASH_REDIS_REST_TOKEN (yoki KV_REST_API_*) - doimiy baza

Alohida Vercel loyihasi (Root Directory: nakladnoy), do'kon botidan mustaqil.
Webhookni ulash: https://<sayt>/api/index?ulash=1&key=<NAKLADNOY_WEBHOOK_SECRET>
"""
from flask import Flask, request, jsonify
import telebot
from telebot.types import InlineKeyboardMarkup, InlineKeyboardButton
from html import escape
from copy import copy
import json, os, io, re, hmac, time, logging, threading, requests

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("nakladnoy")

BOT_TOKEN = os.environ.get("NAKLADNOY_BOT_TOKEN", "")
WEBHOOK_SECRET = os.environ.get("NAKLADNOY_WEBHOOK_SECRET", "")
ADMIN_IDS = {int(x) for x in re.findall(r"\d+", os.environ.get("NAKLADNOY_ADMIN_IDS", "") or os.environ.get("ADMIN_ID", ""))}
# shu guruhlarning barcha a'zolari botdan foydalana oladi (ixtiyoriy): -100123,-100456
GROUP_IDS = {int(x) for x in re.findall(r"-?\d+", os.environ.get("NAKLADNOY_GROUP_IDS", ""))}
REDIS_URL = (os.environ.get("UPSTASH_REDIS_REST_URL") or os.environ.get("KV_REST_API_URL") or "").rstrip("/")
REDIS_TOKEN = os.environ.get("UPSTASH_REDIS_REST_TOKEN") or os.environ.get("KV_REST_API_TOKEN") or ""
SHABLON = os.path.join(os.path.dirname(os.path.abspath(__file__)), "shablon")

MAX_LINES = 1500
MAX_BUTTONS = 10

# Model oldidagi kod: A + brend harfi + tur harfi (ASI = Samsung Incell)
BRANDS = {"S": "Samsung", "I": "iPhone", "R": "Redmi/Xiaomi", "V": "Vivo/Oppo/Realme",
          "C": "Tecno/Infinix", "H": "Honor/Huawei"}
TYPES = {"I": "Incell", "G": "Incell (originalga yaqin)", "T": "TFT", "O": "OLED",
         "J": "Yuqori sifatli nusxa", "R": "Original", "F": "Original"}


def h(v):
    return escape(str(v if v is not None else ""), quote=False)

def som(n):
    return f"{int(round(n or 0)):,}".replace(",", " ")


# ===================== SAQLASH (Upstash Redis, bo'lmasa /tmp) =====================
def _redis(*cmd):
    r = requests.post(REDIS_URL, headers={"Authorization": f"Bearer {REDIS_TOKEN}"}, json=list(cmd), timeout=8)
    r.raise_for_status()
    return r.json().get("result")

def _path(key):
    os.makedirs("/tmp/nakladnoy", exist_ok=True)
    return "/tmp/nakladnoy/" + re.sub(r"[^\w.-]", "_", key) + ".json"

def kv_get(key):
    try:
        if REDIS_URL:
            raw = _redis("GET", "nk:" + key)
            return json.loads(raw) if raw else None
        with open(_path(key)) as f:
            return json.load(f)
    except (FileNotFoundError, ValueError):
        return None
    except Exception:
        log.exception("kv_get %s", key)
        return None

def kv_set(key, value):
    if REDIS_URL:
        _redis("SET", "nk:" + key, json.dumps(value, ensure_ascii=False))
    else:
        with open(_path(key), "w") as f:
            json.dump(value, f, ensure_ascii=False)

def kv_del(key):
    try:
        if REDIS_URL:
            _redis("DEL", "nk:" + key)
        elif os.path.exists(_path(key)):
            os.remove(_path(key))
    except Exception:
        log.exception("kv_del %s", key)

def lines_push(cid, lines):
    """Ro'yxat qatorlarini atomik qo'shadi (bir vaqtda kelgan bir nechta xabar yo'qolmasin)."""
    key = f"lines:{skey(cid)}"
    if REDIS_URL:
        n = 0
        for i in range(0, len(lines), 200):
            n = int(_redis("RPUSH", "nk:" + key, *lines[i:i + 200]))
        return n
    old = kv_get(key) or []
    old += lines
    kv_set(key, old)
    return len(old)

def lines_pop(cid):
    key = f"lines:{skey(cid)}"
    if REDIS_URL:
        out = _redis("LRANGE", "nk:" + key, 0, -1) or []
        _redis("DEL", "nk:" + key)
        return out
    out = kv_get(key) or []
    kv_del(key)
    return out

_ctx = threading.local()   # joriy xabar egasi (guruhda har kimning sessiyasi alohida)

def skey(cid):
    """Shaxsiy chatda: chat ID. Guruhda: chat ID + foydalanuvchi ID."""
    cid = int(cid)
    uid = getattr(_ctx, "uid", None)
    return str(cid) if cid > 0 or not uid else f"{cid}:{uid}"

def get_sess(cid):
    s = kv_get(f"sess:{skey(cid)}")
    return s if isinstance(s, dict) else {}

def set_sess(cid, s):
    kv_set(f"sess:{skey(cid)}", s)


# ===================== BAZA: Xitoy omborining narx fayli =====================
def _rows_from_file(data, filename):
    name = (filename or "").lower()
    if name.endswith(".xls"):
        import xlrd
        sh = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        return [sh.row_values(r) for r in range(sh.nrows)]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith((".txt", ".csv")):
        return [[ln] for ln in data.decode("utf-8", "ignore").splitlines()]
    return None

def _num(v):
    try:
        return int(round(float(str(v).replace(" ", "").replace(",", "")))) if v not in (None, "") else 0
    except ValueError:
        return 0

def _cell(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v if v is not None else "").strip()

CODE_CYR = str.maketrans("АВСЕНКМОРТХУавсенкмортху", "ABCEHKMOPTXYABCEHKMOPTXY")

def norm_code(v):
    """Seriya (kod)ni solishtirish uchun: katta harf, kirill->lotin, faqat harf va raqam, oxiridagi .0 siz."""
    t = _cell(v).upper().translate(CODE_CYR)
    t = re.sub(r"\.0+$", "", t)
    return re.sub(r"[^0-9A-Z]", "", t)

def all_items(catalog):
    """Ekranlar + faqat seriya bo'yicha topiladigan boshqa tovarlar."""
    return (catalog or {}).get("items", []) + (catalog or {}).get("extra", [])

def parse_catalog(rows):
    """Narx faylini o'qiydi. Bu narx fayli bo'lmasa None qaytaradi."""
    for hi, row in enumerate(rows[:15]):
        cells = [str(c or "").strip() for c in row]
        if "存货编码" not in cells:
            continue
        c_code = cells.index("存货编码")
        c_name = cells.index("存货") if "存货" in cells else c_code + 1
        c_price = next((i for i, c in enumerate(cells) if "单价" in c or "UZS" in c.upper()), None)
        c_stock = next((i for i, c in enumerate(cells) if "现存量" in c or "库存" in c), None)
        c_spec = next((i for i, c in enumerate(cells) if "规格型号" in c), None)   # XW-M23
        if c_price is None:
            return None
        items, extra, skipped = [], [], 0
        for row in rows[hi + 1:]:
            row = list(row) + [None] * 20
            code, nm = _cell(row[c_code]), _cell(row[c_name])
            xw = _cell(row[c_spec]) if c_spec is not None else ""
            if not code or not nm:
                continue
            m = re.match(r"^([A-Za-z]{3})\s*-", nm)
            if "停用" in nm or not m:   # to'xtatilgan modellar, quloqchinlar va h.k.
                skipped += 1           # nom bo'yicha qidirilmaydi, lekin seriya yozilsa topiladi
                extra.append({"c": code, "n": nm, "p": _num(row[c_price]),
                              "s": _num(row[c_stock]) if c_stock is not None else None,
                              "b": "?", "t": "?", "x": xw})
                continue
            pre = m.group(1).upper()
            items.append({"c": code, "n": nm, "p": _num(row[c_price]),
                          "s": _num(row[c_stock]) if c_stock is not None else None,
                          "b": pre[1], "t": pre[2], "x": xw})
        return {"items": items, "extra": extra, "skipped": skipped, "updated": int(time.time())}
    return None

def get_catalog():
    c = kv_get("catalog")
    return c if isinstance(c, dict) and c.get("items") else None


# ===================== MODELNI TOPISH =====================
CYR = str.maketrans("АВСЕНКМОРТХУавсенкмортху", "ABCEHKMOPTXYABCEHKMOPTXY")
SEP = re.compile(r"[^0-9A-Z]+")
HEADS = ("REDMI", "REALME", "HM", "RY")          # nom boshidagi brend (HM13C = Redmi 13C, RY X8 = Honor X8)
BRAND_WORDS = {
    "SAMSUNG": "S", "SAMSUNK": "S", "SAMSUNGA": "S",
    "IPHONE": "I", "AYFON": "I", "IFON": "I", "APPLE": "I",
    "REDMI": "R", "REDME": "R", "XIAOMI": "R", "XIOMI": "R", "POCO": "R",
    "VIVO": "V", "OPPO": "V", "REALME": "V",
    "TECNO": "C", "TEKNO": "C", "INFINIX": "C", "ITEL": "C",
    "HONOR": "H", "HUAWEI": "H", "XONOR": "H",
}
TYPE_WORDS = {
    "INCELL": "IG", "INCEL": "IG", "INSEL": "IG", "INSELL": "IG",
    "TFT": "T", "OLED": "O", "AMOLED": "O",
    "ORIGINAL": "RF", "ORGINAL": "RF", "ORIGINALL": "RF", "ORIG": "RF", "ORG": "RF",
    "KOPIYA": "J", "COPY": "J", "NUSXA": "J",
}
FRAME_WORDS = {"RAMKA", "RAMKALI", "RAMKALIK", "PAMKA", "PAMKAJI"}   # kirill "РАМКА" CYR orqali PAMKA bo'ladi
PREFIX_RE = re.compile(r"^A[SIRVCH][IGTOJRF]$")

def chunks(s):
    return [c for c in SEP.split(s.upper().translate(CYR)) if c]

def strip_head(ch):
    if not ch:
        return ch
    first = ch[0]
    if first in BRAND_WORDS or first in ("HM", "RY"):
        return ch[1:]
    for hd in HEADS:
        if first.startswith(hd) and len(first) > len(hd):
            return [first[len(hd):]] + ch[1:]
    return ch

def build_index(items):
    idx = {}
    for it in items:
        body = it["n"].split("-", 1)[1] if "-" in it["n"] else it["n"]
        for alias in body.split("/"):
            ch = chunks(alias)
            for variant in (ch, strip_head(ch)):
                for k in range(1, len(variant) + 1):
                    idx.setdefault("".join(variant[:k]), set()).add(it["c"])
    return idx

QTY_UNIT = re.compile(r"(\d{1,5})\s*(?:ta|dona|шт|sht|pcs|pc|x|х|та)\b\.?", re.I)
QTY_X = re.compile(r"(?:^|\s)[xх×*]\s*(\d{1,5})\b", re.I)
QTY_TAIL = re.compile(r"[\s\-–—:=,]+(\d{1,5})\s*$")

def interpretations(line):
    """Qatordan (model matni, soni) variantlari: avval ehtimoli kattasi."""
    s = re.sub(r"^\s*(?:\d{1,3}[.)]\s+|[-•*·]+\s*)", "", line).strip()
    out = []
    for rx in (QTY_UNIT, QTY_X):
        m = rx.search(s)
        if m:
            out.append(((s[:m.start()] + " " + s[m.end():]).strip(), int(m.group(1))))
    m = QTY_TAIL.search(s)
    if m and not (2010 <= int(m.group(1)) <= 2035) and s[:m.start()].strip():
        out.append((s[:m.start()].strip(), int(m.group(1))))
    out.append((s, None))
    return out

def code_hits(text, codes):
    """Matndagi seriya(lar)ni topadi: butun matn, bo'sh joy bilan ajratilgan so'zlar va ularning juftligi."""
    hits = []
    toks = [t for t in re.split(r"[\s,;]+", text.strip()) if t]
    cands = [text] + toks + [a + b for a, b in zip(toks, toks[1:])] + chunks(text)
    for t in cands:
        k = norm_code(t)
        if len(k) >= 4 and not (k.isdigit() and len(k) < 5) and k in codes:
            hits += [c for c in codes[k] if c not in hits]
    return hits

def lookup(text, cat, idx, codes=None):
    if codes:
        hits = code_hits(text, codes)
        if hits:
            return hits[:1] if len(hits) == 1 else hits
    ch = chunks(text)
    brands, types, prefixes, rest = set(), set(), set(), []
    for c in ch:
        if c in BRAND_WORDS:
            brands.add(BRAND_WORDS[c])
            if c == "POCO":
                rest.append(c)      # POCO model nomining bir qismi ham
        elif c in TYPE_WORDS:
            types.update(TYPE_WORDS[c])
        elif c in FRAME_WORDS:
            rest.append("WF")
        elif PREFIX_RE.match(c):
            prefixes.add(c)
        else:
            rest.append(c)
    if not rest:
        return []
    codes = set()
    for c in rest:                      # to'g'ridan-to'g'ri kod yozilgan bo'lsa (A20106)
        if c in cat:
            codes.add(c)
    if not codes:
        codes = set(idx.get("".join(rest), ()))
        if not codes:
            codes = set(idx.get("".join(strip_head(rest)), ()))
    found = [cat[c] for c in codes if c in cat]
    for keep in (lambda it: it["n"][:3].upper() in prefixes if prefixes else True,
                 lambda it: it["b"] in brands if brands else True,
                 lambda it: it["t"] in types if types else True):
        narrowed = [it for it in found if keep(it)]
        if narrowed:
            found = narrowed
    found.sort(key=lambda it: ("IGTOJRF".find(it["t"]), it["p"]))
    return [it["c"] for it in found]

def resolve_lines(lines, catalog):
    cat = {it["c"].upper(): it for it in all_items(catalog)}
    codes = {}
    for it in all_items(catalog):                        # seriya: A21376
        codes.setdefault(norm_code(it["c"]), []).append(it["c"].upper())
    for it in all_items(catalog):                        # model kodi: XW-M23 (bir nechta bo'lishi mumkin)
        k = norm_code(it.get("x"))
        if k and it["c"].upper() not in codes.setdefault(k, []):
            codes[k].append(it["c"].upper())
    idx = build_index(catalog["items"])
    out = []
    for raw in lines:
        line = raw.strip()
        if not line or not re.search(r"\d", line):
            continue
        item = {"q": line[:120], "qty": 1, "guess": True, "cands": [], "st": "miss"}
        for text, qty in interpretations(line):
            cands = lookup(text, cat, idx, codes)
            if cands:
                item.update(cands=cands[:MAX_BUTTONS], qty=qty or 1, guess=qty is None,
                            st="ok" if len(cands) == 1 else "ask", code=cands[0] if len(cands) == 1 else None)
                break
        orig = {c.upper(): c for c in (x["c"] for x in all_items(catalog))}
        item["cands"] = [orig.get(c, c) for c in item["cands"]]
        if item.get("code"):
            item["code"] = orig.get(item["code"], item["code"])
        out.append(item)
    return out

def split_lines(text):
    out = []
    for ln in (text or "").splitlines():
        parts = [p for p in re.split(r"[;]", ln)]
        for p in parts:
            sub = [x for x in p.split(",")]
            out.extend(sub if len(sub) > 1 and all(re.search(r"\d", x) for x in sub) else [p])
    return [x.strip() for x in out if x.strip()][:MAX_LINES]

def rows_to_lines(rows):
    lines = []
    for r in rows:
        cells = []
        for v in r:
            if v in (None, ""):
                continue
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            cells.append(str(v).strip())
        if len(cells) >= 3 and cells[0].isdigit() and len(cells[0]) <= 4:   # birinchi ustun tartib raqami (№), uzun son esa seriya
            cells = cells[1:]
        if cells:
            lines.append(" ".join(cells))
    return lines[:MAX_LINES]


# ===================== EXCEL: Format1 va Format2 =====================
def _tash_now():
    return time.gmtime(time.time() + 5 * 3600)   # Toshkent vaqti

def _row_style(ws, r, ncol):
    return [copy(ws.cell(r, j)._style) for j in range(1, ncol + 1)], ws.row_dimensions[r].height

def _put_style(ws, r, style):
    cells, height = style
    for j, st in enumerate(cells, start=1):
        ws.cell(r, j)._style = copy(st)
    if height:
        ws.row_dimensions[r].height = height

def merged_items(items, catalog):
    cat = {it["c"]: it for it in all_items(catalog)}
    out = {}
    for it in items:
        if it.get("st") == "ok" and it.get("code") in cat:
            if it["code"] in out:
                out[it["code"]]["qty"] += it["qty"]
            else:
                out[it["code"]] = {**cat[it["code"]], "qty": it["qty"]}
    return list(out.values())

def build_format1(rows, code, tm):
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(SHABLON, "format1.xlsx"))
    ws = wb.active
    data_st, total_st = _row_style(ws, 2, 5), _row_style(ws, 3, 5)
    r = 1
    for it in rows:
        r += 1
        _put_style(ws, r, data_st)
        for j, v in enumerate((it["c"], it["n"], it["qty"], it["p"], f"=C{r}*D{r}"), start=1):
            ws.cell(r, j).value = v
    last = r
    r += 1
    _put_style(ws, r, total_st)
    for j, v in enumerate((code, time.strftime("%d,%m,%Y", tm), f"=SUM(C2:C{last})", None, f"=SUM(E2:E{last})"), start=1):
        ws.cell(r, j).value = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def build_format2(rows, code, tm):
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(SHABLON, "format2.xlsx"))
    ws = wb.active
    data_st, total_st = _row_style(ws, 5, 16), _row_style(ws, 6, 16)
    d = time.strftime("%Y/%m/%d", tm)
    ws["A2"] = f"销售日期：{d}（Savdo sanasi: {d}）"
    first = r = 5
    for it in rows:
        _put_style(ws, r, data_st)
        ws.cell(r, 6).value = it["n"]
        ws.cell(r, 7).value = it["c"]
        ws.cell(r, 9).value = it["p"]
        ws.cell(r, 10).value = it["qty"]
        ws.cell(r, 11).value = f"=J{r}*I{r}"
        r += 1
    last = r - 1
    ws.cell(first, 1).value = code.split("-")[0]
    ws.cell(first, 2).value = code
    if last > first:
        for col in "ABCD":
            ws.merge_cells(f"{col}{first}:{col}{last}")
    _put_style(ws, r, total_st)
    ws.cell(r, 8).value = "TOTAL:"
    ws.cell(r, 10).value = f"=SUM(J{first}:J{last})"
    ws.cell(r, 11).value = f"=SUM(K{first}:K{last})"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ===================== BOT =====================
bot = telebot.TeleBot(BOT_TOKEN or "0:token-sozlanmagan", threaded=False)  # token yo'q bo'lsa ham sayt yiqilmasin
app = Flask(__name__)

def is_admin(uid, chat_id=None):
    return int(uid) in ADMIN_IDS or (chat_id is not None and int(chat_id) in GROUP_IDS)

def is_group(m):
    return m.chat.type in ("group", "supergroup")

def format_kb():
    kb = InlineKeyboardMarkup(row_width=2)
    kb.add(InlineKeyboardButton("📄 Format1", callback_data="nk:f:1"),
           InlineKeyboardButton("📊 Format2", callback_data="nk:f:2"))
    return kb

def collect_kb():
    kb = InlineKeyboardMarkup(row_width=2)
    kb.add(InlineKeyboardButton("✅ Ro'yxat tugadi", callback_data="nk:done"),
           InlineKeyboardButton("❌ Bekor", callback_data="nk:cancel"))
    return kb

FORMAT_DESC = {"1": "Format1 (SERIA | MODEL | SONI | UZS)", "2": "Format2 (TW ombor savdo va to'lov jadvali)"}

def send(cid, text, **kw):
    try:
        return bot.send_message(cid, text, parse_mode="HTML", **kw)
    except Exception:
        log.exception("send -> %s", cid)

def edit(cid, mid, text, kb=None):
    try:
        bot.edit_message_text(text, cid, mid, parse_mode="HTML", reply_markup=kb)
        return True
    except Exception as e:
        if "message is not modified" in str(e):
            return True
        log.warning("edit: %s", e)
        return False

def catalog_line(cat):
    t = time.strftime("%d.%m.%Y %H:%M", time.gmtime(cat["updated"] + 5 * 3600))
    return f"📚 Baza: <b>{len(cat['items'])} ta model</b> (yangilangan: {t})"

def show_menu(cid, prefix=""):
    cat = get_catalog()
    txt = prefix + "🧾 <b>Nakladnoy bot</b>\n\n"
    if cat:
        txt += catalog_line(cat) + "\n\nFormatni tanlang:"
        send(cid, txt, reply_markup=format_kb())
    else:
        send(cid, txt + "⚠️ Hali baza yo'q. Xitoy omborining narx faylini (.xls / .xlsx) shu yerga tashlang.")

def start_format(cid, fmt):
    if not get_catalog():
        return send(cid, "⚠️ Avval Xitoy omborining narx faylini (.xls / .xlsx) tashlang.")
    lines_pop(cid)
    set_sess(cid, {"fmt": fmt, "st": "collect"})
    send(cid, f"✅ <b>{FORMAT_DESC[fmt]}</b>\n\nMijoz ro'yxatini yuboring: matn yoki Excel fayl. "
              "Uzun bo'lsa, bir nechta xabar qilib tashlashingiz mumkin.\n"
              "Masalan:\n<code>A10S 10 ta\nredmi 13c 20\nA02S incell 5</code>\n\n"
              "Hammasini tashlab bo'lgach <b>«✅ Ro'yxat tugadi»</b> ni bosing.", reply_markup=collect_kb())

def ask_text(s, i, cat):
    it = s["items"][i]
    total = s.get("ask_total") or 1
    left = sum(1 for x in s["items"] if x["st"] == "ask")
    pos = max(1, total - left + 1)
    txt = (f"❓ <b>Tanlang ({pos}/{total})</b>\n\nMijoz yozgan: <code>{h(it['q'])}</code>\n"
           f"Soni: <b>{it['qty']}</b>" + (" <i>(yozilmagan, 1 deb olindi)</i>" if it.get("guess") else "") + "\n")
    kb = InlineKeyboardMarkup(row_width=1)
    for k, c in enumerate(it["cands"], start=1):
        p = cat.get(c)
        if not p:
            continue
        tp = TYPES.get(p["t"], "boshqa")
        stock = f" · qoldiq {p['s']}" if p.get("s") is not None else ""
        txt += f"\n<b>{k}.</b> {h(p['n'])}\n     {tp} · <b>{som(p['p'])}</b> so'm{stock}\n"
        kb.add(InlineKeyboardButton(f"{k} · {tp} · {som(p['p'])}", callback_data=f"nk:p:{i}:{c}"))
    kb.add(InlineKeyboardButton("⏭ O'tkazib yuborish", callback_data=f"nk:p:{i}:-"))
    return txt, kb

def next_step(cid, s, mid=None):
    """Navbatdagi ishni bajaradi: tanlash -> topilmaganlar -> kod so'rash."""
    cat = {it["c"]: it for it in all_items(get_catalog())}
    items = s["items"]
    ask = next((k for k, x in enumerate(items) if x["st"] == "ask"), None)
    if ask is not None and not mid:
        s["ask_total"] = sum(1 for x in items if x["st"] == "ask")   # yangi savollar to'plami
    if ask is not None:
        s["st"] = "pick"
        set_sess(cid, s)
        txt, kb = ask_text(s, ask, cat)
        if not (mid and edit(cid, mid, txt, kb)):
            send(cid, txt, reply_markup=kb)
        return
    if mid:
        edit(cid, mid, "✅ Hamma modellar tanlandi.")
    miss = [x for x in items if x["st"] == "miss"]
    if miss:
        s["st"] = "fix"
        set_sess(cid, s)
        lst = "".join(f"• <code>{h(x['q'])}</code>\n" for x in miss[:40])
        more = f"…va yana {len(miss) - 40} ta\n" if len(miss) > 40 else ""
        kb = InlineKeyboardMarkup().add(InlineKeyboardButton("➡️ Ularsiz davom etish", callback_data="nk:cont"))
        send(cid, f"❌ <b>Bazadan topilmadi ({len(miss)} ta):</b>\n{lst}{more}\n"
                  "To'g'rilab qayta yozib yuboring (masalan modelni aniqroq yozing), "
                  "yoki ularsiz davom eting.", reply_markup=kb)
        return
    ask_code(cid, s)

def ask_code(cid, s):
    cat = get_catalog()
    rows = merged_items(s["items"], cat)
    if not rows:
        set_sess(cid, {})
        return show_menu(cid, "⚠️ Ro'yxatda birorta model qolmadi.\n\n")
    s["st"] = "code"
    set_sess(cid, s)
    total = sum(r["p"] * r["qty"] for r in rows)
    txt = (f"📦 <b>Tayyor:</b> {len(rows)} model, <b>{sum(r['qty'] for r in rows)} dona</b>, "
           f"<b>{som(total)} so'm</b>\n")
    short = [r for r in rows if r.get("s") is not None and r["qty"] > r["s"]]
    if short:
        txt += "\n⚠️ <b>Omborda yetmaydi:</b>\n" + "".join(
            f"• {h(r['n'][:40])} — so'raldi {r['qty']}, bor {r['s']}\n" for r in short[:20])
    guess = [x for x in s["items"] if x["st"] == "ok" and x.get("guess")]
    if guess:
        txt += "\n⚠️ <b>Soni yozilmagan, 1 deb olindi:</b>\n" + "".join(f"• <code>{h(x['q'])}</code>\n" for x in guess[:20])
    skipped = sum(1 for x in s["items"] if x["st"] in ("skip", "miss"))
    if skipped:
        txt += f"\n⏭ Kirmay qolgan qatorlar: {skipped} ta\n"
    txt += "\n✍️ <b>Qaysi kod bilan?</b> (masalan <code>WE-029</code>)"
    kb = InlineKeyboardMarkup().add(InlineKeyboardButton("❌ Bekor", callback_data="nk:cancel"))
    send(cid, txt, reply_markup=kb)

def other_kb(fmt):
    alt = "2" if fmt == "1" else "1"
    icon = "📊" if alt == "2" else "📄"
    return InlineKeyboardMarkup().add(InlineKeyboardButton(f"{icon} Shuni Format{alt}'da ham", callback_data=f"nk:alt:{alt}"))

def send_file(cid, rows, code, fmt, tm):
    data = (build_format1 if fmt == "1" else build_format2)(rows, code, tm)
    fname = f"{code} {time.strftime('%d.%m.%Y', tm)}.xlsx"
    total = sum(r["p"] * r["qty"] for r in rows)
    bot.send_document(cid, io.BytesIO(data), visible_file_name=fname, parse_mode="HTML",
                      caption=f"🧾 <b>{h(code)}</b> · Format{fmt}\n{len(rows)} model, "
                              f"{sum(r['qty'] for r in rows)} dona, {som(total)} so'm",
                      reply_markup=other_kb(fmt))

def finish(cid, s, code):
    rows = merged_items(s["items"], get_catalog())
    tm = _tash_now()
    send_file(cid, rows, code, s["fmt"], tm)
    # oxirgi nakladnoy: boshqa formatda ham olish uchun (ro'yxatni qayta kiritmasdan)
    kv_set(f"last:{skey(cid)}", {"rows": rows, "code": code, "tm": list(tm)[:9]})
    set_sess(cid, {})
    send(cid, "Yana nakladnoy kerak bo'lsa, formatni tanlang:", reply_markup=format_kb())


# ---------- buyruqlar ----------
def admin_msg(handler):
    def wrapper(m):
        _ctx.uid = m.from_user.id
        if not is_admin(m.from_user.id, m.chat.id):
            log.info("Admin emas, yozdi: id=%s @%s chat=%s", m.from_user.id, m.from_user.username, m.chat.id)
            if is_group(m):
                return            # guruhda begonalarga javob bermaymiz (spam bo'lmasin)
            return send(m.chat.id, f"⛔️ Bu bot faqat admin uchun.\nSizning ID: <code>{m.from_user.id}</code>")
        return handler(m)
    wrapper.__name__ = handler.__name__
    return wrapper

def group_hint(m):
    """Guruhda bot oddiy xabarlarni ko'ra oladimi (privacy mode)."""
    if not is_group(m):
        return ""
    try:
        me = bot.get_me()
        st = bot.get_chat_member(m.chat.id, me.id).status
        if me.can_read_all_group_messages or st in ("administrator", "creator"):
            return ""
    except Exception:
        log.exception("group_hint")
        return ""
    return ("⚠️ <b>Guruhda ro'yxatni ko'rishim uchun meni guruh admini qiling</b> "
            "(yoki ro'yxatni mening xabarimga <i>javob (reply)</i> qilib yuboring).\n\n")

@bot.message_handler(commands=["start", "menu"])
@admin_msg
def cmd_start(m):
    set_sess(m.chat.id, {})
    show_menu(m.chat.id, group_hint(m))

@bot.message_handler(commands=["format1", "format2"])
@admin_msg
def cmd_format(m):
    start_format(m.chat.id, "1" if m.text.lower().startswith("/format1") else "2")

@bot.message_handler(commands=["bekor"])
@admin_msg
def cmd_cancel(m):
    lines_pop(m.chat.id)
    set_sess(m.chat.id, {})
    show_menu(m.chat.id, "❌ Bekor qilindi.\n\n")

@bot.message_handler(commands=["baza"])
@admin_msg
def cmd_baza(m):
    cat = get_catalog()
    if not cat:
        return send(m.chat.id, "⚠️ Hali baza yo'q. Narx faylini (.xls / .xlsx) tashlang.")
    by = {}
    for it in cat["items"]:
        k = f"A{it['b']}{it['t']}"
        by[k] = by.get(k, 0) + 1
    lst = "".join(f"• <b>{k}</b> — {BRANDS.get(k[1], '?')}, {TYPES.get(k[2], '?')}: {n}\n"
                  for k, n in sorted(by.items(), key=lambda x: -x[1]))
    send(m.chat.id, catalog_line(cat) + f"\n\n{lst}\nYangilash uchun yangi narx faylini tashlang.")


@bot.message_handler(content_types=["document"])
@admin_msg
def on_document(m):
    cid, doc = m.chat.id, m.document
    if doc.file_size and doc.file_size > 15 * 1024 * 1024:
        return send(cid, "⚠️ Fayl juda katta.")
    try:
        data = bot.download_file(bot.get_file(doc.file_id).file_path)
        rows = _rows_from_file(data, doc.file_name)
    except Exception:
        log.exception("fayl o'qish")
        return send(cid, "⚠️ Faylni o'qib bo'lmadi. .xls, .xlsx yoki .txt bo'lsin.")
    if rows is None:
        return send(cid, "⚠️ Faqat .xls, .xlsx yoki .txt fayl qabul qilinadi.")
    cat = parse_catalog(rows)
    if cat:
        if not cat["items"]:
            return send(cid, "⚠️ Narx faylida birorta model topilmadi.")
        kv_set("catalog", cat)
        types = {}
        for it in cat["items"]:
            types[TYPES.get(it["t"], it["t"])] = types.get(TYPES.get(it["t"], it["t"]), 0) + 1
        lst = ", ".join(f"{k} {v}" for k, v in sorted(types.items(), key=lambda x: -x[1]))
        txt = f"✅ <b>Baza yangilandi:</b> {len(cat['items'])} ta model\n{lst}"
        if not REDIS_URL:
            txt += "\n\n⚠️ Doimiy baza (Redis) ulanmagan: fayl vaqtincha saqlanadi, keyin qayta tashlash kerak bo'lishi mumkin."
        if cat["skipped"]:
            txt += f"\n⏭ Tashlab ketildi (停用 / ekran emas): {cat['skipped']} ta"
        s = get_sess(cid)
        if s.get("st") == "collect":
            return send(cid, txt + "\n\nRo'yxatni yuborishda davom eting.", reply_markup=collect_kb())
        return show_menu(cid, txt + "\n\n")
    s = get_sess(cid)
    if s.get("st") != "collect":
        if is_group(m):
            return
        return show_menu(cid, "ℹ️ Bu narx fayli emas. Ro'yxat bo'lsa, avval formatni tanlang.\n\n")
    add_lines(cid, rows_to_lines(rows))

def add_lines(cid, lines):
    if not lines:
        return send(cid, "⚠️ Bu xabarda ro'yxat topilmadi.", reply_markup=collect_kb())
    n = lines_push(cid, lines)
    send(cid, f"📥 +{len(lines)} qator (jami <b>{n}</b>). Yana yuboring yoki <b>«✅ Ro'yxat tugadi»</b> ni bosing.",
         reply_markup=collect_kb())

@bot.message_handler(content_types=["text"])
@admin_msg
def on_text(m):
    cid, s = m.chat.id, get_sess(m.chat.id)
    st = s.get("st")
    if st == "collect":
        return add_lines(cid, split_lines(m.text))
    if st == "fix":
        cat = get_catalog()
        new = resolve_lines(split_lines(m.text), cat)
        if not new:
            return send(cid, "⚠️ Model topilmadi. Qayta yozing yoki «➡️ Ularsiz davom etish» ni bosing.")
        for x in s["items"]:
            if x["st"] == "miss":
                x["st"] = "skip"
        s["items"] += new
        return next_step(cid, s)
    if st == "pick":
        return None if is_group(m) else send(cid, "☝️ Avval yuqoridagi savolga tugma orqali javob bering (yoki /bekor).")
    if st == "code":
        code = re.sub(r"\s+", "", m.text.strip().upper())
        if not re.fullmatch(r"[A-Z]{1,5}-\d{1,5}", code):
            return send(cid, "⚠️ Kod <code>WE-029</code> ko'rinishida bo'lsin. Qayta yozing:")
        try:
            return finish(cid, s, code)
        except Exception:
            log.exception("finish")
            return send(cid, "⚠️ Fayl tayyorlashda xatolik. Qayta urinib ko'ring.")
    if not is_group(m):     # guruhdagi oddiy suhbatga javob bermaymiz
        show_menu(cid)


# ---------- tugmalar ----------
@bot.callback_query_handler(func=lambda c: (c.data or "").startswith("nk:"))
def on_callback(c):
    cid = c.message.chat.id
    _ctx.uid = c.from_user.id
    if not is_admin(c.from_user.id, cid):
        return bot.answer_callback_query(c.id, "Faqat admin uchun.", show_alert=True)
    parts = c.data.split(":")
    act = parts[1]
    try:
        bot.answer_callback_query(c.id)
    except Exception:
        pass
    if act == "alt" and parts[2] in ("1", "2"):
        last = kv_get(f"last:{skey(cid)}")
        if not last or not last.get("rows"):
            return send(cid, "⚠️ Oxirgi nakladnoy topilmadi. Ro'yxatni qaytadan yuboring.", reply_markup=format_kb())
        try:
            return send_file(cid, last["rows"], last["code"], parts[2], time.struct_time(tuple(last["tm"])))
        except Exception:
            log.exception("alt format")
            return send(cid, "⚠️ Fayl tayyorlashda xatolik. Qayta urinib ko'ring.")
    if act == "f" and parts[2] in ("1", "2"):
        return start_format(cid, parts[2])
    if act == "cancel":
        lines_pop(cid)
        set_sess(cid, {})
        return show_menu(cid, "❌ Bekor qilindi.\n\n")
    s = get_sess(cid)
    if act == "done":
        if s.get("st") != "collect":
            return
        lines = lines_pop(cid)
        if not lines:
            return send(cid, "⚠️ Hali ro'yxat yuborilmadi.", reply_markup=collect_kb())
        cat = get_catalog()
        items = resolve_lines(lines, cat)
        s["items"] = items
        ok = sum(1 for x in items if x["st"] == "ok")
        ask = sum(1 for x in items if x["st"] == "ask")
        miss = sum(1 for x in items if x["st"] == "miss")
        send(cid, f"📋 <b>{len(items)} qator:</b> ✅ {ok} topildi · ❓ {ask} tanlash kerak · ❌ {miss} topilmadi")
        return next_step(cid, s)
    if act == "p" and s.get("st") == "pick" and len(parts) == 4:
        i = int(parts[2])
        if i >= len(s["items"]) or s["items"][i]["st"] != "ask":
            return
        if parts[3] == "-":
            s["items"][i]["st"] = "skip"
        elif parts[3] in s["items"][i]["cands"]:
            s["items"][i].update(st="ok", code=parts[3])
        return next_step(cid, s, mid=c.message.message_id)
    if act == "cont" and s.get("st") == "fix":
        for x in s["items"]:
            if x["st"] == "miss":
                x["st"] = "skip"
        try:
            bot.edit_message_reply_markup(cid, c.message.message_id, reply_markup=None)
        except Exception:
            pass
        return ask_code(cid, s)


# ===================== WEBHOOK =====================
def do_ulash():
    key = request.args.get("key", "")
    if not WEBHOOK_SECRET or not hmac.compare_digest(key, WEBHOOK_SECRET):
        return "Forbidden: kalit noto'g'ri", 403
    try:
        url = f"https://{request.host}/api/index"
        ok = bot.set_webhook(url=url, secret_token=WEBHOOK_SECRET, allowed_updates=["message", "callback_query"])
        me = bot.get_me()
        return (f"✅ @{h(me.username)} {h(url)} manziliga ulandi. Telegramda /start yozing.", 200) if ok \
            else ("❌ Webhook ulanmadi", 500)
    except Exception as e:
        log.exception("set_webhook")
        return f"XATOLIK: {h(e)}", 500

@app.route("/", defaults={"path": ""}, methods=["POST", "GET"])
@app.route("/<path:path>", methods=["POST", "GET"])
def webhook(path):
    if request.method == "GET":
        if "ulash" in request.args:
            return do_ulash()
        return ("✅ Nakladnoy bot ishlayapti." if BOT_TOKEN else "⚠️ NAKLADNOY_BOT_TOKEN sozlanmagan."), 200
    token = request.headers.get("X-Telegram-Bot-Api-Secret-Token", "")
    if not WEBHOOK_SECRET or not hmac.compare_digest(token, WEBHOOK_SECRET):
        return "Forbidden", 403
    try:
        bot.process_new_updates([telebot.types.Update.de_json(request.get_data().decode("utf-8"))])
    except Exception:
        log.exception("process update")
    return jsonify({"status": "ok"}), 200
